from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "DOCUMENTACION_TECNICA_BASE_DE_DATOS_SYNCUT.md"
TARGET = ROOT / "docs" / "DOCUMENTACION_TECNICA_BASE_DE_DATOS_SYNCUT.docx"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def add_inline(paragraph, text):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(31, 78, 121)
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def parse_table(lines, start, document):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    if not rows:
        return index
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for col, value in enumerate(rows[0]):
        cell = table.rows[0].cells[col]
        cell.text = value
        shade(cell, "1F4E79")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for source_row in rows[1:]:
        target_row = table.add_row().cells
        for col, value in enumerate(source_row):
            target_row[col].text = value
    document.add_paragraph()
    return index


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    for name, size, color in (("Title", 24, "1F4E79"), ("Heading 1", 16, "1F4E79"), ("Heading 2", 13, "2F5597"), ("Heading 3", 11, "4472C4")):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("SyncUT — Documentación técnica de base de datos | Uso interno")

    in_code = False
    code_lines = []
    index = 0
    title_done = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                paragraph = document.add_paragraph()
                paragraph.style = document.styles["Normal"]
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                shade_dummy = OxmlElement("w:shd")
                shade_dummy.set(qn("w:fill"), "F2F2F2")
                paragraph._p.get_or_add_pPr().append(shade_dummy)
                in_code = False
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if stripped.startswith("|"):
            index = parse_table(lines, index, document)
            continue
        if stripped in ("", "---"):
            index += 1
            continue
        if line.startswith("# "):
            if title_done:
                document.add_heading(line[2:], 1)
            else:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run(line[2:])
                title_done = True
        elif line.startswith("## "):
            document.add_heading(line[3:], 1)
        elif line.startswith("### "):
            document.add_heading(line[4:], 2)
        elif line.startswith("#### "):
            document.add_heading(line[5:], 3)
        elif re.match(r"^\d+\. ", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\. ", "", stripped))
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, stripped.replace("  ", " "))
        index += 1

    document.core_properties.title = "SyncUT — Documentación formal y técnica de la base de datos"
    document.core_properties.subject = "Arquitectura, modelo, seguridad y operación de PostgreSQL/Supabase"
    document.core_properties.author = "Equipo SyncUT"
    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
